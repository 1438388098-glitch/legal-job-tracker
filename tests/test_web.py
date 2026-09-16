import pytest
from fastapi.testclient import TestClient

from app.web.main import create_app


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_DISABLE_SCHEDULER", "1")
    monkeypatch.setattr("app.snapshot.SNAP_DIR", tmp_path / "snap")
    app = create_app(str(tmp_path / "w.db"))
    app.state.conn.execute(
        "INSERT INTO sources(slug,name,kind,config) VALUES('demo','演示','html','{}')")
    app.state.conn.execute(
        "INSERT INTO jobs(title,org,job_type,city,deadline,source_slug,url,url_fingerprint,"
        "search_text) VALUES('书记员招聘','某法院','public','深圳', date('now','+2 day'),"
        "'demo','https://d/1','fp1','书记员招聘 某法院 报名2026年10月1日截止')")
    app.state.conn.commit()
    with TestClient(app) as c:
        yield c


def test_list_and_filters(client):
    r = client.get("/")
    assert r.status_code == 200 and "书记员招聘" in r.text
    r = client.get("/", params={"q": "报名"})
    assert "书记员招聘" in r.text
    r = client.get("/", params={"city": "广州"})
    assert "书记员招聘" not in r.text


def test_detail_and_status(client):
    r = client.get("/jobs/1")
    assert r.status_code == 200 and "书记员招聘" in r.text
    r = client.post("/jobs/1/status", data={"status": "read"})
    assert r.status_code in (200, 303)
    st = client.app.state.conn.execute(
        "SELECT status FROM jobs WHERE id=1").fetchone()["status"]
    assert st == "read"


def _add(conn, title, url, kind):
    conn.execute(
        "INSERT INTO jobs(title,source_slug,url,url_fingerprint,search_text,notice_kind) "
        "VALUES(?,?,'https://d/'||?,'fp'||?,?,?)", (title, "demo", url, url, title, kind))
    conn.commit()


def test_result_notices_hidden_by_default(client):
    """事后结果公示默认不出现在列表里，但可搜索、可切换查看。"""
    conn = client.app.state.conn
    _add(conn, "某单位2026年公开招聘拟聘用人员名单公示", "r1", "result")
    _add(conn, "某单位2026年公开招聘工作人员公告", "o1", "opening")

    assert "拟聘用人员名单公示" not in client.get("/").text
    assert "公开招聘工作人员公告" in client.get("/").text
    assert "拟聘用人员名单公示" in client.get("/", params={"notice_kind": "result"}).text
    # 搜索仍能召回（用户想查自己那次考试结果时用得上）。
    # 用 2 字关键词走 LIKE 分支：测试直接插表，没建 FTS 索引，3 字以上会走 FTS 查不到。
    assert "拟聘用人员名单公示" in client.get(
        "/", params={"q": "拟聘", "notice_kind": "all"}).text


def test_export_defaults_to_openings(client):
    conn = client.app.state.conn
    _add(conn, "某某公司招聘法务专员公告", "o2", "opening")
    _add(conn, "某某公司拟录用人员名单公示", "r2", "result")

    import io

    import openpyxl
    r = client.get("/export.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    titles = [row[1] for row in ws.iter_rows(values_only=True)][1:]
    assert any("法务专员" in (t or "") for t in titles)
    assert not any("拟录用人员名单公示" in (t or "") for t in titles)

    ws2 = openpyxl.load_workbook(io.BytesIO(
        client.get("/export.xlsx", params={"notice_kind": "all"}).content)).active
    titles2 = [row[1] for row in ws2.iter_rows(values_only=True)][1:]
    assert any("拟录用人员名单公示" in (t or "") for t in titles2)


def test_apply_flow(client):
    client.post("/jobs/1/apply")
    r = client.get("/board")
    assert "待投" in r.text
    aid = client.app.state.conn.execute("SELECT id FROM applications").fetchone()["id"]
    client.post(f"/applications/{aid}/status", data={"status": "已投", "note": "官网投递"})
    st = client.app.state.conn.execute(
        "SELECT status FROM applications WHERE id=?", (aid,)).fetchone()["status"]
    assert st == "已投"


def test_paste_flow(client, monkeypatch):
    monkeypatch.setattr("app.collect.pastebox.collect_url", lambda url: {
        "title": "律所实习招聘", "url": url, "source_slug": "pastebox", "org": None,
        "city": "深圳", "job_type": "intern", "publish_date": None,
        "deadline": None, "body": "正文", "status": "pending", "needs_review": True})
    r = client.post("/paste", data={"url": "https://mp.weixin.qq.com/s/x"})
    assert r.status_code in (200, 303)
    jid = client.app.state.conn.execute(
        "SELECT id FROM jobs WHERE status='pending'").fetchone()["id"]
    r = client.post(f"/jobs/{jid}/confirm", data={
        "title": "律所实习招聘(修)", "city": "深圳", "deadline": "2026-10-01",
        "job_type": "intern", "notes": ""})
    assert r.status_code in (200, 303)
    st = client.app.state.conn.execute(
        "SELECT status FROM jobs WHERE id=?", (jid,)).fetchone()["status"]
    # confirm 把状态重置为 new；TestClient 会跟随重定向进详情页，
    # 而"看过即已读"会把它标成 read —— 这正是新交互的预期结果
    assert st == "read"


def test_health_page(client):
    r = client.get("/health")
    assert r.status_code == 200 and "演示" in r.text


def test_detail_auto_marks_read(client):
    """看过详情就算已读，不用再点一次"标记已读"。"""
    assert client.app.state.conn.execute(
        "SELECT status FROM jobs WHERE id=1").fetchone()["status"] == "new"
    client.get("/jobs/1")
    assert client.app.state.conn.execute(
        "SELECT status FROM jobs WHERE id=1").fetchone()["status"] == "read"


def test_actions_return_to_filtered_list(client):
    """操作完必须带着筛选参数回来，不能把用户的视图归零。"""
    url = "/?job_type=public&city=%E6%B7%B1%E5%9C%B3"      # 类型/城市筛选
    client.get(url)
    r = client.post("/jobs/1/status", data={"status": "read"},
                    headers={"referer": "http://testserver" + url},
                    follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"].startswith("/?")          # 保留了 query
    assert "job_type=public" in r.headers["location"]
    # 没有来源页时退回默认
    r2 = client.post("/jobs/1/status", data={"status": "new"}, follow_redirects=False)
    assert r2.headers["location"] == "/"


def test_apply_from_list_row(client):
    """列表行直接加跟踪：不用进详情页。"""
    r = client.post("/jobs/1/apply", headers={"referer": "http://testserver/"},
                    follow_redirects=False)
    assert r.status_code == 303
    assert client.app.state.conn.execute(
        "SELECT COUNT(*) c FROM applications").fetchone()["c"] == 1
    # 再点一次不会重复
    client.post("/jobs/1/apply", headers={"referer": "http://testserver/"})
    assert client.app.state.conn.execute(
        "SELECT COUNT(*) c FROM applications").fetchone()["c"] == 1


def test_list_row_shows_apply_state(client):
    r = client.get("/")
    assert "跟踪" in r.text        # 未跟踪 → 显示"跟踪"按钮
    client.post("/jobs/1/apply")
    r = client.get("/")
    assert "取消跟踪" in r.text    # 已跟踪 → 可撤销（原来只能加不能删）


def test_batch_operations(client):
    conn = client.app.state.conn
    _add(conn, "某某公司招聘法务专员公告", "b1", "opening")
    _add(conn, "某某公司招聘法律顾问公告", "b2", "opening")
    ids = ",".join(str(r["id"]) for r in conn.execute("SELECT id FROM jobs"))
    r = client.post("/jobs/batch", data={"action": "read", "ids": ids},
                    follow_redirects=False)
    assert r.status_code == 303
    left = conn.execute(
        "SELECT COUNT(*) c FROM jobs WHERE status='read'").fetchone()["c"]
    assert left == 3
    # 空 ids 不应报错也不应改动
    client.post("/jobs/batch", data={"action": "archive", "ids": ""})
    assert conn.execute(
        "SELECT COUNT(*) c FROM jobs WHERE status='archived'").fetchone()["c"] == 0


def test_synonym_search_widens_recall(client):
    """搜"律所"必须也能命中写了"律师事务所"的记录（同义词组展开）。

    "律所"是 2 字，走 LIKE 分支；同组的"律师事务所"出现在另一条记录里。
    """
    _add(client.app.state.conn, "广东华某律师事务所招聘授薪律师公告", "s1", "opening")
    r = client.get("/", params={"q": "律所"})
    assert "广东华某律师事务所招聘授薪律师公告" in r.text


def test_sort_by_created(client):
    r = client.get("/", params={"sort": "new"})
    assert r.status_code == 200 and "书记员招聘" in r.text


def test_fresh_and_nodeadline_entries(client):
    """概览条上的"今日新增 / 无截止日"必须是入口，不是摆设。"""
    conn = client.app.state.conn
    conn.execute("UPDATE jobs SET created_at='2020-01-01 08:00:00' WHERE id=1")
    conn.commit()
    assert "书记员招聘" not in client.get("/", params={"fresh": "1"}).text
    conn.execute("UPDATE jobs SET deadline=NULL, rolling=1 WHERE id=1")
    conn.commit()
    assert "书记员招聘" in client.get("/", params={"nodeadline": "1"}).text


def test_rolling_shown_as_long_term(client):
    """长期有效岗位没有截止日，但要显示"长期"而不是"—"，否则像缺数据。"""
    conn = client.app.state.conn
    conn.execute("UPDATE jobs SET deadline=NULL, rolling=1 WHERE id=1")
    conn.commit()
    r = client.get("/")
    assert "长期" in r.text and "deadline none" not in r.text


def test_export_xlsx(client):
    r = client.get("/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml")


def test_unapply(client):
    """误点「跟踪」要能撤销：取消后岗位回到未跟踪状态，可重新加入。"""
    client.post("/jobs/1/apply")
    aid = client.app.state.conn.execute(
        "SELECT id FROM applications WHERE job_id=1").fetchone()["id"]
    r = client.post(f"/applications/{aid}/remove", follow_redirects=False)
    assert r.status_code == 303
    assert client.app.state.conn.execute(
        "SELECT COUNT(*) c FROM applications").fetchone()["c"] == 0
    # 撤销后可以重新跟踪
    client.post("/jobs/1/apply")
    assert client.app.state.conn.execute(
        "SELECT COUNT(*) c FROM applications").fetchone()["c"] == 1


def test_unapply(client):
    """误点「跟踪」要能撤销：取消后岗位回到未跟踪状态，可重新加入。"""
    client.post("/jobs/1/apply")
    aid = client.app.state.conn.execute(
        "SELECT id FROM applications WHERE job_id=1").fetchone()["id"]
    r = client.post(f"/applications/{aid}/remove", follow_redirects=False)
    assert r.status_code == 303
    assert client.app.state.conn.execute(
        "SELECT COUNT(*) c FROM applications").fetchone()["c"] == 0
    # 撤销后可以重新跟踪
    client.post("/jobs/1/apply")
    assert client.app.state.conn.execute(
        "SELECT COUNT(*) c FROM applications").fetchone()["c"] == 1
